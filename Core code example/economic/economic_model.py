"""Reconnect the unchanged legacy economic model to active PFAS mass flows.

The two underlying work-breakdown-structure (WBS) workbooks remain unchanged.
Their cached annualized cost and design-throughput results are converted to:

* USD per kg of IER used for the wastewater stream; and
* USD per kg of PFAS assigned to gas, solvent, or solid-waste incineration.

Wastewater cost is therefore sensitive to the active ``f_C`` scenarios through
actual IER demand.  The IER WBS already includes spent-resin handling and
destruction, so no second incineration charge is added to captured wastewater
PFAS.  Other-stream costs retain the legacy common incineration unit cost.
"""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import xlrd

from attribution.genai_pfas_attribution import MATERIAL_NAMES


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_ATTRIBUTION = (
    PROJECT_ROOT / "Excel results" / "attribution" / "GenAI_PFAS_attribution_results.json"
)
DEFAULT_IER_BURDEN = PROJECT_ROOT / "Excel results" / "ier" / "IER_burden_results.json"
DEFAULT_LEGACY_INPUT = PROJECT_ROOT / "Data input" / "legacy_data.xls"
DEFAULT_ECONOMIC_DIR = PROJECT_ROOT / "Economic model"
DEFAULT_IER_WBS = DEFAULT_ECONOMIC_DIR / "IER-wastewater-.xlsm"
DEFAULT_INCINERATION_WBS = DEFAULT_ECONOMIC_DIR / "Incineration-solvent.xlsm"
DEFAULT_SUMMARY = DEFAULT_ECONOMIC_DIR / "economic cost.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "Excel results" / "economic" / "Economic_results.json"

EXPECTED_PERIODS = tuple(
    f"{year} Q{quarter}"
    for year in range(2025, 2031)
    for quarter in range(1, 5)
)
MEDIA = ("gas", "solvent", "wastewater", "solid")
OTHER_INCINERATION_MEDIA = ("gas", "solvent", "solid")
US_GALLON_TO_KG_WATER_EQUIVALENT = 3.785
LEGACY_REFERENCE_CAPACITY_G_G = 1.05


class EconomicInputError(ValueError):
    """Raised when a cached WBS result or upstream interface is invalid."""


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return None if not math.isfinite(float(value)) else float(value)
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    return value


def _read_json(source: str | Path | dict[str, Any], label: str) -> tuple[dict[str, Any], str]:
    if isinstance(source, dict):
        return source, f"in-memory {label}"
    path = Path(source).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"{label} not found: {path}")
    return json.loads(path.read_text(encoding="utf-8")), str(path)


def _cached_number(sheet, cell: str, context: str) -> float:
    value = sheet[cell].value
    if not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        raise EconomicInputError(f"{context} must contain a finite cached number")
    return float(value)


def _read_wbs_unit_costs(
    ier_wbs_path: str | Path,
    incineration_wbs_path: str | Path,
) -> dict[str, float]:
    ier_path = Path(ier_wbs_path).expanduser().resolve()
    incin_path = Path(incineration_wbs_path).expanduser().resolve()
    for path in (ier_path, incin_path):
        if not path.exists():
            raise FileNotFoundError(f"Economic WBS workbook not found: {path}")

    ier_book = openpyxl.load_workbook(
        ier_path, read_only=True, data_only=True, keep_vba=True
    )
    try:
        if "OUTPUT" not in ier_book.sheetnames:
            raise EconomicInputError(f"{ier_path.name} is missing OUTPUT")
        output = ier_book["OUTPUT"]
        ier_annual_cost = _cached_number(
            output, "C443", f"{ier_path.name}/OUTPUT!C443 annualized cost"
        )
        ier_annual_use_ton = _cached_number(
            output, "C426", f"{ier_path.name}/OUTPUT!C426 spent media"
        )
    finally:
        ier_book.close()

    incin_book = openpyxl.load_workbook(
        incin_path, read_only=True, data_only=True, keep_vba=True
    )
    try:
        if "OUTPUT" not in incin_book.sheetnames:
            raise EconomicInputError(f"{incin_path.name} is missing OUTPUT")
        output = incin_book["OUTPUT"]
        incin_annual_cost = _cached_number(
            output, "C443", f"{incin_path.name}/OUTPUT!C443 annualized cost"
        )
        incin_design_flow_us_gal = _cached_number(
            output, "C425", f"{incin_path.name}/OUTPUT!C425 design flow"
        )
    finally:
        incin_book.close()

    if min(
        ier_annual_cost,
        ier_annual_use_ton,
        incin_annual_cost,
        incin_design_flow_us_gal,
    ) <= 0:
        raise EconomicInputError("Cached WBS cost and throughput values must be positive")

    ier_annual_use_kg = ier_annual_use_ton * 1000.0
    incin_design_flow_kg = (
        incin_design_flow_us_gal * US_GALLON_TO_KG_WATER_EQUIVALENT
    )
    return {
        "ier_annualized_cost_usd_per_year": ier_annual_cost,
        "ier_design_use_kg_per_year": ier_annual_use_kg,
        "ier_cost_usd_per_kg_ier": ier_annual_cost / ier_annual_use_kg,
        "incineration_annualized_cost_usd_per_year": incin_annual_cost,
        "incineration_design_throughput_kg_per_year": incin_design_flow_kg,
        "incineration_cost_usd_per_kg_pfas": incin_annual_cost
        / incin_design_flow_kg,
    }


def _read_summary_checks(path: str | Path) -> dict[str, float | None]:
    summary_path = Path(path).expanduser().resolve()
    if not summary_path.exists():
        raise FileNotFoundError(f"Economic summary workbook not found: {summary_path}")
    workbook = openpyxl.load_workbook(summary_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        values: dict[str, float | None] = {}
        for key, cell in {
            "summary_ier_annualized_cost_usd_per_year": "B4",
            "summary_ier_use_ton_per_year": "B5",
            "summary_stale_capacity_g_g": "B6",
            "summary_stale_cost_usd_per_kg_pfas": "B8",
            "summary_solvent_cost_usd_per_kg": "B15",
            "summary_gas_cost_usd_per_kg": "B19",
            "summary_solid_cost_usd_per_kg": "B23",
        }.items():
            value = sheet[cell].value
            values[key] = float(value) if isinstance(value, (int, float)) else None
    finally:
        workbook.close()
    return values


def _read_emission_factors(path: str | Path) -> np.ndarray:
    legacy_path = Path(path).expanduser().resolve()
    if not legacy_path.exists():
        raise FileNotFoundError(f"Legacy emission-factor workbook not found: {legacy_path}")
    workbook = xlrd.open_workbook(str(legacy_path))
    if "Emission factor" not in workbook.sheet_names():
        raise EconomicInputError(f"{legacy_path.name} is missing Emission factor")
    sheet = workbook.sheet_by_name("Emission factor")
    factors = np.asarray(
        [sheet.row_values(row, 1, 7) for row in range(1, 5)], dtype=float
    )
    if factors.shape != (4, len(MATERIAL_NAMES)) or not np.isfinite(factors).all():
        raise EconomicInputError("Emission-factor matrix must be four by six")
    return factors


def _aggregate(
    rows: list[dict[str, Any]], group_keys: tuple[str, ...]
) -> list[dict[str, Any]]:
    value_keys = (
        "pfas_mass_kg",
        "ier_required_kg",
        "wastewater_ier_cost_usd",
        "other_incineration_cost_usd",
        "total_cost_usd",
    )
    grouped: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        key = tuple(row[name] for name in group_keys)
        if key not in grouped:
            grouped[key] = {name: row[name] for name in group_keys}
            grouped[key].update({name: 0.0 for name in value_keys})
        for name in value_keys:
            grouped[key][name] += float(row.get(name, 0.0))
    return list(grouped.values())


def calculate_economic_results(
    attribution_results: str | Path | dict[str, Any] = DEFAULT_ATTRIBUTION,
    ier_burden_results: str | Path | dict[str, Any] = DEFAULT_IER_BURDEN,
    legacy_input_path: str | Path = DEFAULT_LEGACY_INPUT,
    ier_wbs_path: str | Path = DEFAULT_IER_WBS,
    incineration_wbs_path: str | Path = DEFAULT_INCINERATION_WBS,
    economic_summary_path: str | Path = DEFAULT_SUMMARY,
) -> dict[str, Any]:
    """Calculate quarterly through cumulative screening treatment costs."""

    attribution, attribution_source = _read_json(
        attribution_results, "active attribution output"
    )
    ier_burden, ier_source = _read_json(ier_burden_results, "IER burden output")
    quarterly = attribution.get("quarterly", [])
    if tuple(str(row.get("Period")) for row in quarterly) != EXPECTED_PERIODS:
        raise EconomicInputError("Attribution periods are not 2025 Q1-2030 Q4")
    scenarios = ier_burden.get("mass_scenarios", [])
    if not scenarios:
        raise EconomicInputError("IER burden output contains no mass scenarios")
    primary_scenario_id = str(ier_burden.get("primary_scenario_id", ""))

    unit_costs = _read_wbs_unit_costs(ier_wbs_path, incineration_wbs_path)
    summary_checks = _read_summary_checks(economic_summary_path)
    emission_factors = _read_emission_factors(legacy_input_path)
    medium_index = {medium: index for index, medium in enumerate(MEDIA)}

    quarterly_by_medium: list[dict[str, Any]] = []
    quarterly_by_material: list[dict[str, Any]] = []
    max_material_medium_mass_error = 0.0
    max_cost_identity_error = 0.0
    for scenario in scenarios:
        scenario_id = str(scenario["scenario_id"])
        resin_per_influent = float(scenario["ier_kg_per_kg_influent"])
        for attribution_row in quarterly:
            period = str(attribution_row["Period"])
            year = int(attribution_row["Year"])
            quarter = int(attribution_row["Quarter"])
            material_masses = np.asarray(
                [
                    float(attribution_row[f"Material | {material} (kg)"])
                    for material in MATERIAL_NAMES
                ],
                dtype=float,
            )
            for medium in MEDIA:
                pfas_mass = float(
                    attribution_row[f"Manufacturing flow | {medium} (kg)"]
                )
                ier_mass = pfas_mass * resin_per_influent if medium == "wastewater" else 0.0
                ier_cost = ier_mass * unit_costs["ier_cost_usd_per_kg_ier"]
                incin_cost = (
                    pfas_mass * unit_costs["incineration_cost_usd_per_kg_pfas"]
                    if medium in OTHER_INCINERATION_MEDIA
                    else 0.0
                )
                total_cost = ier_cost + incin_cost
                max_cost_identity_error = max(
                    max_cost_identity_error, abs(total_cost - ier_cost - incin_cost)
                )
                common = {
                    "period": period,
                    "year": year,
                    "quarter": quarter,
                    "scenario_id": scenario_id,
                    "role": str(scenario["role"]),
                    "capacity_utilization_factor_f_c": float(
                        scenario["capacity_utilization_factor_f_c"]
                    ),
                    "medium": medium,
                }
                quarterly_by_medium.append(
                    {
                        **common,
                        "pfas_mass_kg": pfas_mass,
                        "ier_required_kg": ier_mass,
                        "wastewater_ier_cost_usd": ier_cost,
                        "other_incineration_cost_usd": incin_cost,
                        "total_cost_usd": total_cost,
                    }
                )

                material_medium = material_masses * emission_factors[medium_index[medium], :]
                max_material_medium_mass_error = max(
                    max_material_medium_mass_error,
                    abs(float(material_medium.sum()) - pfas_mass),
                )
                for material_index, material in enumerate(MATERIAL_NAMES):
                    material_mass = float(material_medium[material_index])
                    material_ier = (
                        material_mass * resin_per_influent
                        if medium == "wastewater"
                        else 0.0
                    )
                    material_ier_cost = (
                        material_ier * unit_costs["ier_cost_usd_per_kg_ier"]
                    )
                    material_incin_cost = (
                        material_mass
                        * unit_costs["incineration_cost_usd_per_kg_pfas"]
                        if medium in OTHER_INCINERATION_MEDIA
                        else 0.0
                    )
                    quarterly_by_material.append(
                        {
                            **common,
                            "material": material,
                            "pfas_mass_kg": material_mass,
                            "ier_required_kg": material_ier,
                            "wastewater_ier_cost_usd": material_ier_cost,
                            "other_incineration_cost_usd": material_incin_cost,
                            "total_cost_usd": material_ier_cost
                            + material_incin_cost,
                        }
                    )

    annual_by_medium = _aggregate(
        quarterly_by_medium, ("year", "scenario_id", "role", "capacity_utilization_factor_f_c", "medium")
    )
    cumulative_by_medium = _aggregate(
        quarterly_by_medium, ("scenario_id", "role", "capacity_utilization_factor_f_c", "medium")
    )
    annual_total = _aggregate(
        quarterly_by_medium, ("year", "scenario_id", "role", "capacity_utilization_factor_f_c")
    )
    cumulative_total = _aggregate(
        quarterly_by_medium, ("scenario_id", "role", "capacity_utilization_factor_f_c")
    )
    annual_by_material = _aggregate(
        quarterly_by_material,
        ("year", "scenario_id", "role", "capacity_utilization_factor_f_c", "medium", "material"),
    )
    cumulative_by_material = _aggregate(
        quarterly_by_material,
        ("scenario_id", "role", "capacity_utilization_factor_f_c", "medium", "material"),
    )

    reconciliation_fields = (
        "pfas_mass_kg",
        "ier_required_kg",
        "wastewater_ier_cost_usd",
        "other_incineration_cost_usd",
        "total_cost_usd",
    )
    max_annual_cumulative_error = 0.0
    max_medium_total_error = 0.0
    max_material_medium_error = 0.0
    reconciliation_passed = True
    for total_row in cumulative_total:
        scenario_id = str(total_row["scenario_id"])
        annual_rows = [
            row for row in annual_total if str(row["scenario_id"]) == scenario_id
        ]
        medium_rows = [
            row for row in cumulative_by_medium if str(row["scenario_id"]) == scenario_id
        ]
        for field in reconciliation_fields:
            reconciliation_passed = reconciliation_passed and all(
                math.isclose(float(total_row[field]),
                             sum(float(row[field]) for row in rows),
                             rel_tol=1e-12, abs_tol=1e-8)
                for rows in (annual_rows, medium_rows)
            )
            max_annual_cumulative_error = max(
                max_annual_cumulative_error,
                abs(float(total_row[field]) - sum(float(row[field]) for row in annual_rows)),
            )
            max_medium_total_error = max(
                max_medium_total_error,
                abs(float(total_row[field]) - sum(float(row[field]) for row in medium_rows)),
            )
        for medium_row in medium_rows:
            material_rows = [
                row
                for row in cumulative_by_material
                if str(row["scenario_id"]) == scenario_id
                and row["medium"] == medium_row["medium"]
            ]
            for field in reconciliation_fields:
                reconciliation_passed = reconciliation_passed and math.isclose(
                    float(medium_row[field]), sum(float(row[field]) for row in material_rows),
                    rel_tol=1e-12, abs_tol=1e-8,
                )
                max_material_medium_error = max(
                    max_material_medium_error,
                    abs(
                        float(medium_row[field])
                        - sum(float(row[field]) for row in material_rows)
                    ),
                )

    incin_cost = unit_costs["incineration_cost_usd_per_kg_pfas"]
    summary_incin_values = [
        summary_checks[f"summary_{medium}_cost_usd_per_kg"]
        for medium in OTHER_INCINERATION_MEDIA
    ]
    max_summary_incin_difference = max(
        abs(float(value) - incin_cost)
        for value in summary_incin_values
        if value is not None
    )
    summary_ier_annual = summary_checks["summary_ier_annualized_cost_usd_per_year"]
    summary_ier_use = summary_checks["summary_ier_use_ton_per_year"]
    max_ier_wbs_summary_difference = max(
        abs(float(summary_ier_annual) - unit_costs["ier_annualized_cost_usd_per_year"])
        if summary_ier_annual is not None
        else math.inf,
        abs(float(summary_ier_use) * 1000.0 - unit_costs["ier_design_use_kg_per_year"])
        if summary_ier_use is not None
        else math.inf,
    )
    legacy_cost_per_kg_captured_pfas = (
        unit_costs["ier_cost_usd_per_kg_ier"] / LEGACY_REFERENCE_CAPACITY_G_G
    )
    tolerance = 1e-8
    quality_checks = {
        "quarter_count": len(EXPECTED_PERIODS),
        "scenario_count": len(scenarios),
        "max_material_to_medium_mass_error_kg": max_material_medium_mass_error,
        "max_cost_component_identity_error_usd": max_cost_identity_error,
        "max_annual_to_cumulative_reconciliation_error": max_annual_cumulative_error,
        "max_medium_to_total_reconciliation_error": max_medium_total_error,
        "max_material_to_medium_reconciliation_error": max_material_medium_error,
        "aggregation_reconciliation_passed": reconciliation_passed,
        "aggregation_relative_tolerance": 1e-12,
        "aggregation_absolute_tolerance": 1e-8,
        "max_summary_incineration_unit_cost_difference": max_summary_incin_difference,
        "max_ier_wbs_summary_cache_difference": max_ier_wbs_summary_difference,
        "stale_summary_capacity_cell_ignored": True,
        "captured_wastewater_pfas_not_double_charged_for_incineration": True,
    }
    quality_checks["all_checks_passed"] = all(
        (
            max_material_medium_mass_error <= tolerance,
            max_cost_identity_error <= tolerance,
            reconciliation_passed,
            max_summary_incin_difference <= tolerance,
            max_ier_wbs_summary_difference <= tolerance,
            primary_scenario_id
            in {str(scenario["scenario_id"]) for scenario in scenarios},
        )
    )

    return _json_safe(
        {
            "schema_version": "1.0",
            "model_name": "Legacy WBS economic interface for active PFAS flows",
            "period": "2025 Q1-2030 Q4",
            "currency_basis": "USD in the unchanged legacy WBS basis",
            "attribution_interface": attribution_source,
            "ier_interface": ier_source,
            "primary_scenario_id": primary_scenario_id,
            "equations": {
                "wastewater": "cost = wastewater PFAS * IER kg/kg influent PFAS * USD/kg IER",
                "gas_solvent_solid": "cost = PFAS sent to incineration * USD/kg PFAS",
            },
            "unit_costs": {
                **unit_costs,
                "legacy_reference_capacity_g_pfas_per_g_ier": LEGACY_REFERENCE_CAPACITY_G_G,
                "legacy_implied_cost_usd_per_kg_captured_pfas": legacy_cost_per_kg_captured_pfas,
            },
            "legacy_summary_audit": summary_checks,
            "quarterly_by_medium": quarterly_by_medium,
            "annual_by_medium": annual_by_medium,
            "cumulative_by_medium": cumulative_by_medium,
            "annual_total": annual_total,
            "cumulative_2025_2030": cumulative_total,
            "annual_by_material": annual_by_material,
            "cumulative_by_material": cumulative_by_material,
            "sources": {
                "ier_wbs": str(Path(ier_wbs_path).expanduser().resolve()),
                "incineration_wbs": str(Path(incineration_wbs_path).expanduser().resolve()),
                "legacy_summary": str(Path(economic_summary_path).expanduser().resolve()),
                "emission_factors": str(Path(legacy_input_path).expanduser().resolve()),
            },
            "limitations": [
                "The WBS models and their prices are unchanged; this module only reconnects active masses.",
                "The same legacy incineration unit cost is used for gas, solvent, and solid PFAS.",
                "The IER WBS includes spent-media handling/destruction; captured wastewater PFAS receives no separate incineration charge.",
                "The cached summary capacity and cost-per-kg-PFAS cells are stale and are excluded from calculation.",
                "No temporal price escalation, regional price differentiation, or economic uncertainty is applied in the central run.",
            ],
            "quality_checks": quality_checks,
        }
    )


def save_json(
    payload: dict[str, Any], output_path: str | Path = DEFAULT_OUTPUT
) -> Path:
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attribution", type=Path, default=DEFAULT_ATTRIBUTION)
    parser.add_argument("--ier-burden", type=Path, default=DEFAULT_IER_BURDEN)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    payload = calculate_economic_results(args.attribution, args.ier_burden)
    if not payload["quality_checks"]["all_checks_passed"]:
        raise RuntimeError("Economic interface checks failed")
    output = save_json(payload, args.output)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

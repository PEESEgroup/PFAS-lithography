"""Rebuild IER treatment burdens from transparent mass-dependent components.

The legacy IER factor combines production/separation of one kilogram of resin
with incineration of the resulting spent-resin mixture.  This module separates
those two contributions before applying the capacity-utilization sensitivity:

    BF_influent = m_IER * S + (m_IER + m_captured) * D + m_captured * L

where ``S`` is impact per kilogram of IER produced, ``D`` is impact per
kilogram of spent IER mixture excluding lime, and L is lime burden per kg
captured PFAS. Every mass is expressed per
kilogram of PFAS entering wastewater management.  The 90% capture assumption
therefore changes both the captured-PFAS mass and the spent-mixture mass.  The
unmodified legacy factor is recovered exactly on a per-kilogram-captured basis
when effective capacity is 1.05 g/g (historical calibration only).
"""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
import sys
from typing import Any, Iterable

import numpy as np
import openpyxl
import xlrd

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ier.ier_model import DEFAULT_INPUT as DEFAULT_IER_INPUT
from ier.ier_model import run_ier_model


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_ATTRIBUTION = (
    PROJECT_ROOT / "Excel results" / "attribution" / "GenAI_PFAS_attribution_results.json"
)
DEFAULT_LEGACY_FACTOR = (
    PROJECT_ROOT / "Data input" / "treatment" / "treatment_destruction_summary_new.xls"
)
DEFAULT_LCI_2025 = (
    PROJECT_ROOT
    / "Data input"
    / "treatment"
    / "treatment_Database-IAM_scenarios_Basic_2025.xlsx"
)
DEFAULT_LCI_2030 = (
    PROJECT_ROOT / "Data input" / "treatment" / "treatment_PFAS_destruction_NDI_2030.xlsx",
    PROJECT_ROOT
    / "Data input"
    / "treatment"
    / "treatment_PFAS_destruction_PkDuge500_2030.xlsx",
    PROJECT_ROOT
    / "Data input"
    / "treatment"
    / "treatment_PFAS_destruction_PkDuge1150_2030.xlsx",
    PROJECT_ROOT
    / "Data input"
    / "treatment"
    / "treatment_PFAS_destruction_SSP5-BASE_2030.xlsx",
)
DEFAULT_OUTPUT = PROJECT_ROOT / "Excel results" / "ier" / "IER_burden_results.json"

IMPACT_CODES = (
    "CC",
    "AC",
    "ECF_total",
    "ECF_inorganics",
    "ECF_organics",
    "ER_nonrenewable",
    "EF",
    "EM",
    "ET",
    "HTC",
    "HTN",
    "IR",
    "LU",
    "MR",
    "OD",
    "PMF",
    "POF",
    "WU",
    "CED",
)
IMPACT_COUNT = len(IMPACT_CODES)
REFERENCE_CAPTURED_PFAS_KG = 1.05
REFERENCE_IER_KG = 1.0
REFERENCE_SPENT_MIX_KG = REFERENCE_IER_KG + REFERENCE_CAPTURED_PFAS_KG
START_YEAR = 2025
END_YEAR = 2030
EXPECTED_QUARTERS = (END_YEAR - START_YEAR + 1) * 4


class IERBurdenInputError(ValueError):
    """Raised when an IER burden source violates the expected input contract."""


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return None if not math.isfinite(float(value)) else float(value)
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    return value


def _as_numeric_vector(values: Iterable[Any], context: str) -> np.ndarray:
    try:
        vector = np.asarray(list(values), dtype=float)
    except (TypeError, ValueError) as exc:
        raise IERBurdenInputError(f"{context} contains non-numeric values") from exc
    if vector.shape != (IMPACT_COUNT,):
        raise IERBurdenInputError(
            f"{context} must contain {IMPACT_COUNT} impacts; found {vector.shape}"
        )
    if not np.isfinite(vector).all():
        raise IERBurdenInputError(f"{context} contains non-finite values")
    return vector


def _read_lci_rows(path: Path) -> tuple[np.ndarray, np.ndarray, list[str], np.ndarray, np.ndarray]:
    if not path.exists():
        raise FileNotFoundError(f"IER LCI workbook not found: {path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "IER" not in workbook.sheetnames:
            raise IERBurdenInputError(f"{path.name} is missing sheet 'IER'")
        sheet = workbook["IER"]
        separation_label = str(sheet.cell(33, 2).value or "").strip().lower()
        destruction_label = str(sheet.cell(34, 2).value or "").strip().lower()
        if "separation" not in separation_label:
            raise IERBurdenInputError(
                f"{path.name} row 33 is not the IER separation/production result"
            )
        if "destruction" not in destruction_label:
            raise IERBurdenInputError(
                f"{path.name} row 34 is not the spent-mixture destruction result"
            )
        separation = _as_numeric_vector(
            (sheet.cell(33, column).value for column in range(5, 24)),
            f"{path.name} IER row 33",
        )
        destruction_batch = _as_numeric_vector(
            (sheet.cell(34, column).value for column in range(5, 24)),
            f"{path.name} IER row 34",
        )
        method_names = [
            str(sheet.cell(2, column).value or "").strip()
            for column in range(5, 24)
        ]
        component_sheet = workbook['IER_factor']
        if not (math.isclose(float(sheet['D3'].value), REFERENCE_IER_KG)
                and math.isclose(float(sheet['A37'].value), REFERENCE_CAPTURED_PFAS_KG)):
            raise IERBurdenInputError(f'{path.name}: historical LCI batch basis changed')
        component_separation = _as_numeric_vector(
            (component_sheet.cell(14, c).value for c in range(5, 24)),
            f'{path.name} IER_factor row 14',
        )
        if not np.allclose(component_separation, separation, rtol=1e-10, atol=1e-12):
            raise IERBurdenInputError(f'{path.name}: resin factor does not match source inventory')
        separation = component_separation
        destruction_no_lime = _as_numeric_vector(
            (component_sheet.cell(15, c).value for c in range(5, 24)),
            f'{path.name} IER_factor row 15',
        )
        lime_per_pfas = _as_numeric_vector(
            (component_sheet.cell(16, c).value for c in range(5, 24)),
            f'{path.name} IER_factor row 16',
        )
        if not np.allclose(destruction_no_lime * REFERENCE_SPENT_MIX_KG
                           + lime_per_pfas * REFERENCE_CAPTURED_PFAS_KG,
                           destruction_batch, rtol=1e-10, atol=1e-12):
            raise IERBurdenInputError(f'{path.name}: component factors do not reconstruct the legacy batch')
    finally:
        workbook.close()
    return separation, destruction_batch, method_names, destruction_no_lime, lime_per_pfas


def _read_legacy_reference(
    path: Path,
) -> tuple[list[str], list[str], list[dict[str, Any]], np.ndarray]:
    if not path.exists():
        raise FileNotFoundError(f"Legacy IER factor workbook not found: {path}")
    workbook = xlrd.open_workbook(str(path))
    if "data for python-IER" not in workbook.sheet_names():
        raise IERBurdenInputError(f"{path.name} is missing 'data for python-IER'")
    if "2030" not in workbook.sheet_names():
        raise IERBurdenInputError(f"{path.name} is missing '2030'")

    factor_sheet = workbook.sheet_by_name("data for python-IER")
    unit_sheet = workbook.sheet_by_name("2030")
    legacy_headers = [str(value).strip() for value in factor_sheet.row_values(0, 1, 20)]
    units = [str(value).strip() for value in unit_sheet.row_values(1, 1, 20)]
    if len(legacy_headers) != IMPACT_COUNT or len(units) != IMPACT_COUNT:
        raise IERBurdenInputError("Legacy factor workbook does not contain 19 impacts")

    rows: list[dict[str, Any]] = []
    values: list[list[float]] = []
    for row_index in range(1, factor_sheet.nrows):
        decimal_year = factor_sheet.cell_value(row_index, 0)
        if not isinstance(decimal_year, (int, float)):
            continue
        if START_YEAR - 1e-9 <= decimal_year <= END_YEAR + 0.75 + 1e-9:
            year = int(math.floor(decimal_year + 1e-9))
            quarter = int(round((decimal_year - year) * 4)) + 1
            rows.append(
                {
                    "period": f"{year} Q{quarter}",
                    "year": year,
                    "quarter": quarter,
                    "decimal_year": float(decimal_year),
                }
            )
            values.append(factor_sheet.row_values(row_index, 1, 20))
    factor_array = np.asarray(values, dtype=float)
    if factor_array.shape != (EXPECTED_QUARTERS, IMPACT_COUNT):
        raise IERBurdenInputError(
            "Legacy IER central-factor block must contain 24 quarters by 19 impacts; "
            f"found {factor_array.shape}"
        )
    return legacy_headers, units, rows, factor_array


def _read_attribution(source: str | Path | dict[str, Any]) -> dict[str, Any]:
    if isinstance(source, dict):
        return source
    path = Path(source).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"Attribution output not found: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def _expected_periods() -> list[str]:
    return [
        f"{year} Q{quarter}"
        for year in range(START_YEAR, END_YEAR + 1)
        for quarter in range(1, 5)
    ]


def calculate_ier_burden(
    attribution_results: str | Path | dict[str, Any] = DEFAULT_ATTRIBUTION,
    ier_model_results: dict[str, Any] | None = None,
    ier_input_path: str | Path = DEFAULT_IER_INPUT,
    legacy_factor_path: str | Path = DEFAULT_LEGACY_FACTOR,
    lci_2025_path: str | Path = DEFAULT_LCI_2025,
    lci_2030_paths: Iterable[str | Path] = DEFAULT_LCI_2030,
    *, require_legacy_regression: bool = True,
) -> dict[str, Any]:
    """Calculate quarterly, annual, and cumulative IER treatment burdens."""

    attribution_source = (
        "in-memory output from active GenAI attribution"
        if isinstance(attribution_results, dict)
        else str(Path(attribution_results).expanduser().resolve())
    )
    attribution = _read_attribution(attribution_results)
    if ier_model_results is None:
        ier_model_results = run_ier_model(ier_input_path)

    legacy_path = Path(legacy_factor_path).expanduser().resolve()
    source_2025_path = Path(lci_2025_path).expanduser().resolve()
    source_2030_paths = [Path(path).expanduser().resolve() for path in lci_2030_paths]
    if len(source_2030_paths) != 4:
        raise IERBurdenInputError("Exactly four 2030 IAM scenario workbooks are required")

    separation_2025, destruction_2025_batch, method_names_2025, destruction_no_lime_2025, lime_2025 = _read_lci_rows(
        source_2025_path
    )
    scenario_2030 = [_read_lci_rows(path) for path in source_2030_paths]
    separation_2030 = np.mean(
        np.stack([values[0] for values in scenario_2030]), axis=0
    )
    destruction_2030_batch = np.mean(
        np.stack([values[1] for values in scenario_2030]), axis=0
    )
    method_names = scenario_2030[0][2]
    destruction_no_lime_2030 = np.mean(np.stack([v[3] for v in scenario_2030]), axis=0)
    lime_2030 = np.mean(np.stack([v[4] for v in scenario_2030]), axis=0)
    if any(values[2] != method_names for values in scenario_2030[1:]):
        raise IERBurdenInputError("The four 2030 workbooks use different impact methods")
    if len(method_names_2025) != IMPACT_COUNT:
        raise IERBurdenInputError("The 2025 workbook does not expose 19 impact methods")

    legacy_headers, units, legacy_periods, legacy_factors = _read_legacy_reference(
        legacy_path
    )
    expected_periods = _expected_periods()
    if [row["period"] for row in legacy_periods] != expected_periods:
        raise IERBurdenInputError("Legacy IER factor periods are not 2025 Q1-2030 Q4")

    quarterly_attribution = attribution.get("quarterly", [])
    if len(quarterly_attribution) != EXPECTED_QUARTERS:
        raise IERBurdenInputError(
            f"Attribution must contain {EXPECTED_QUARTERS} quarters; "
            f"found {len(quarterly_attribution)}"
        )
    if [str(row.get("Period")) for row in quarterly_attribution] != expected_periods:
        raise IERBurdenInputError("Attribution periods are not 2025 Q1-2030 Q4")

    scenario_rows = ier_model_results.get("sensitivity_results", [])
    if not scenario_rows:
        raise IERBurdenInputError("IER mass model produced no enabled scenarios")
    primary_id = str(ier_model_results.get("primary_scenario_id", ""))

    impacts = []
    for index, code in enumerate(IMPACT_CODES):
        unit = units[index]
        if not unit and code in {"ECF_inorganics", "ECF_organics"}:
            unit = "CTUe"
        elif not unit and code == "ER_nonrenewable":
            unit = "source unit not retained in legacy summary"
        impacts.append(
            {
                "impact_index": index + 1,
                "impact_code": code,
                "legacy_header": legacy_headers[index],
                "method_name": method_names[index],
                "unit": unit,
            }
        )

    interpolated_separation = []
    interpolated_destruction_batch = []
    for quarter_index in range(EXPECTED_QUARTERS):
        fraction = quarter_index / (EXPECTED_QUARTERS - 1)
        interpolated_separation.append(
            separation_2025 + (separation_2030 - separation_2025) * fraction
        )
        interpolated_destruction_batch.append(
            destruction_2025_batch
            + (destruction_2030_batch - destruction_2025_batch) * fraction
        )
    separation_factors = np.stack(interpolated_separation)
    destruction_batch_factors = np.stack(interpolated_destruction_batch)
    fractions = np.linspace(0.0, 1.0, EXPECTED_QUARTERS)[:, None]
    destruction_per_mix_kg = destruction_no_lime_2025 + fractions * (destruction_no_lime_2030 - destruction_no_lime_2025)
    lime_per_captured_kg = lime_2025 + fractions * (lime_2030 - lime_2025)
    reconstructed_legacy = (
        separation_factors + destruction_per_mix_kg * REFERENCE_SPENT_MIX_KG
        + lime_per_captured_kg * REFERENCE_CAPTURED_PFAS_KG
    ) / REFERENCE_CAPTURED_PFAS_KG
    reconstruction_difference = reconstructed_legacy - legacy_factors

    mass_scenarios: list[dict[str, Any]] = []
    quarterly_factors: list[dict[str, Any]] = []
    quarterly_burden: list[dict[str, Any]] = []
    max_total_identity_error = 0.0
    max_total_identity_relative_error = 0.0
    primary_captured_differences: list[float] = []
    primary_influent_differences: list[float] = []

    for scenario in scenario_rows:
        basis = float(scenario["influent_pfas_basis_kg"])
        if basis <= 0:
            raise IERBurdenInputError("Influent PFAS basis must be positive")
        capture_fraction = float(scenario["capture_fraction"])
        resin_per_influent = float(scenario["ier_mass_kg"]) / basis
        captured_per_influent = float(scenario["captured_pfas_kg"]) / basis
        bypass_per_influent = float(scenario["bypass_pfas_to_water_kg"]) / basis
        mix_per_influent = resin_per_influent + captured_per_influent
        scenario_id = str(scenario["scenario_id"])
        role = str(scenario["role"])
        mass_scenarios.append(
            {
                "scenario_id": scenario_id,
                "role": role,
                "capacity_utilization_factor_f_c": float(
                    scenario["capacity_utilization_factor_f_c"]
                ),
                "effective_capacity_g_pfas_per_g_ier": float(
                    scenario["effective_capacity_g_pfas_per_g_ier"]
                ),
                "capture_fraction": capture_fraction,
                "captured_pfas_kg_per_kg_influent": captured_per_influent,
                "bypass_pfas_kg_per_kg_influent": bypass_per_influent,
                "ier_kg_per_kg_influent": resin_per_influent,
                "spent_mix_kg_per_kg_influent": mix_per_influent,
            }
        )

        for quarter_index, attribution_row in enumerate(quarterly_attribution):
            period = expected_periods[quarter_index]
            year = int(attribution_row["Year"])
            quarter = int(attribution_row["Quarter"])
            fraction = quarter_index / (EXPECTED_QUARTERS - 1)
            influent_pfas_kg = float(
                attribution_row["Manufacturing flow | wastewater (kg)"]
            )
            if influent_pfas_kg < 0:
                raise IERBurdenInputError(f"Negative wastewater PFAS in {period}")

            for impact_index, impact in enumerate(impacts):
                separation_per_ier = float(
                    separation_factors[quarter_index, impact_index]
                )
                destruction_per_mix = float(
                    destruction_per_mix_kg[quarter_index, impact_index]
                )
                separation_per_influent = resin_per_influent * separation_per_ier
                lime_per_captured = float(lime_per_captured_kg[quarter_index, impact_index])
                mix_no_lime_per_influent = mix_per_influent * destruction_per_mix
                lime_per_influent = captured_per_influent * lime_per_captured
                destruction_per_influent = mix_no_lime_per_influent + lime_per_influent
                total_per_influent = separation_per_influent + destruction_per_influent
                total_per_captured = total_per_influent / capture_fraction
                legacy_factor = float(legacy_factors[quarter_index, impact_index])

                separation_total = influent_pfas_kg * separation_per_influent
                destruction_total = influent_pfas_kg * destruction_per_influent
                total_burden = influent_pfas_kg * total_per_influent
                identity_error = total_burden - separation_total - destruction_total
                max_total_identity_error = max(
                    max_total_identity_error, abs(identity_error)
                )
                max_total_identity_relative_error = max(
                    max_total_identity_relative_error,
                    abs(identity_error) / max(abs(total_burden), 1.0),
                )

                common = {
                    "period": period,
                    "year": year,
                    "quarter": quarter,
                    "interpolation_fraction": fraction,
                    "scenario_id": scenario_id,
                    "role": role,
                    "capacity_utilization_factor_f_c": float(
                        scenario["capacity_utilization_factor_f_c"]
                    ),
                    "impact_index": impact["impact_index"],
                    "impact_code": impact["impact_code"],
                    "unit": impact["unit"],
                }
                quarterly_factors.append(
                    {
                        **common,
                        "separation_factor_per_kg_ier": separation_per_ier,
                        "destruction_factor_per_reference_2p05kg_batch": float(
                            destruction_batch_factors[quarter_index, impact_index]
                        ),
                        "destruction_excluding_lime_factor_per_kg_spent_mix": destruction_per_mix,
                        "lime_factor_per_kg_captured_pfas": lime_per_captured,
                        "lime_burden_factor_per_kg_influent_pfas": lime_per_influent,
                        "separation_burden_factor_per_kg_influent_pfas": separation_per_influent,
                        "destruction_burden_factor_per_kg_influent_pfas": destruction_per_influent,
                        "total_burden_factor_per_kg_influent_pfas": total_per_influent,
                        "total_burden_factor_per_kg_captured_pfas": total_per_captured,
                        "legacy_total_factor_per_kg_captured_pfas": legacy_factor,
                    }
                )
                quarterly_burden.append(
                    {
                        **common,
                        "influent_wastewater_pfas_kg": influent_pfas_kg,
                        "captured_pfas_kg": influent_pfas_kg * captured_per_influent,
                        "bypass_pfas_to_water_kg": influent_pfas_kg * bypass_per_influent,
                        "ier_required_kg": influent_pfas_kg * resin_per_influent,
                        "spent_mix_incinerated_kg": influent_pfas_kg * mix_per_influent,
                        "separation_burden": separation_total,
                        "spent_mix_incineration_burden": destruction_total,
                        "incineration_excluding_lime_burden": influent_pfas_kg * mix_no_lime_per_influent,
                        "pfas_specific_lime_burden": influent_pfas_kg * lime_per_influent,
                        "total_ier_treatment_burden": total_burden,
                    }
                )
                if scenario_id == primary_id:
                    independent_factor = (
                        separation_per_ier / float(scenario['effective_capacity_g_pfas_per_g_ier'])
                        + destruction_per_mix * (1.0 + 1.0 / float(scenario['effective_capacity_g_pfas_per_g_ier']))
                        + lime_per_captured
                    )
                    primary_captured_differences.append(
                        total_per_captured - independent_factor
                    )
                    primary_influent_differences.append(
                        total_per_influent - capture_fraction * independent_factor
                    )

    annual_burden: list[dict[str, Any]] = []
    cumulative_burden: list[dict[str, Any]] = []
    for scenario in mass_scenarios:
        scenario_id = scenario["scenario_id"]
        for impact in impacts:
            impact_code = impact["impact_code"]
            scenario_impact_rows = [
                row
                for row in quarterly_burden
                if row["scenario_id"] == scenario_id
                and row["impact_code"] == impact_code
            ]
            for year in range(START_YEAR, END_YEAR + 1):
                year_rows = [row for row in scenario_impact_rows if row["year"] == year]
                if len(year_rows) != 4:
                    raise AssertionError(
                        f"{scenario_id}/{impact_code}/{year} does not contain four quarters"
                    )
                annual_burden.append(
                    {
                        "year": year,
                        "scenario_id": scenario_id,
                        "role": scenario["role"],
                        "capacity_utilization_factor_f_c": scenario[
                            "capacity_utilization_factor_f_c"
                        ],
                        "impact_index": impact["impact_index"],
                        "impact_code": impact_code,
                        "unit": impact["unit"],
                        "influent_wastewater_pfas_kg": sum(
                            row["influent_wastewater_pfas_kg"] for row in year_rows
                        ),
                        "captured_pfas_kg": sum(
                            row["captured_pfas_kg"] for row in year_rows
                        ),
                        "bypass_pfas_to_water_kg": sum(
                            row["bypass_pfas_to_water_kg"] for row in year_rows
                        ),
                        "ier_required_kg": sum(
                            row["ier_required_kg"] for row in year_rows
                        ),
                        "spent_mix_incinerated_kg": sum(
                            row["spent_mix_incinerated_kg"] for row in year_rows
                        ),
                        "separation_burden": sum(
                            row["separation_burden"] for row in year_rows
                        ),
                        "spent_mix_incineration_burden": sum(
                            row["spent_mix_incineration_burden"] for row in year_rows
                        ),
                        "total_ier_treatment_burden": sum(
                            row["total_ier_treatment_burden"] for row in year_rows
                        ),
                    }
                )
            cumulative_burden.append(
                {
                    "period": f"{START_YEAR}-{END_YEAR}",
                    "scenario_id": scenario_id,
                    "role": scenario["role"],
                    "capacity_utilization_factor_f_c": scenario[
                        "capacity_utilization_factor_f_c"
                    ],
                    "impact_index": impact["impact_index"],
                    "impact_code": impact_code,
                    "unit": impact["unit"],
                    "separation_burden": sum(
                        row["separation_burden"] for row in scenario_impact_rows
                    ),
                    "spent_mix_incineration_burden": sum(
                        row["spent_mix_incineration_burden"]
                        for row in scenario_impact_rows
                    ),
                    "total_ier_treatment_burden": sum(
                        row["total_ier_treatment_burden"]
                        for row in scenario_impact_rows
                    ),
                }
            )

    max_reconstruction_difference = float(np.max(np.abs(reconstruction_difference)))
    max_primary_captured_difference = max(
        (abs(value) for value in primary_captured_differences), default=math.inf
    )
    max_primary_influent_difference = max(
        (abs(value) for value in primary_influent_differences), default=math.inf
    )
    tolerance = 1e-10
    quality_checks = {
        "quarter_count": EXPECTED_QUARTERS,
        "scenario_count": len(mass_scenarios),
        "impact_count": IMPACT_COUNT,
        "quarterly_factor_row_count": len(quarterly_factors),
        "quarterly_burden_row_count": len(quarterly_burden),
        "annual_burden_row_count": len(annual_burden),
        "cumulative_burden_row_count": len(cumulative_burden),
        "max_legacy_factor_reconstruction_abs_difference": max_reconstruction_difference,
        "legacy_factor_reconstruction_passed": max_reconstruction_difference <= tolerance,
        "legacy_regression_required": require_legacy_regression,
        "max_primary_per_captured_factor_abs_difference_vs_three_component_formula": max_primary_captured_difference,
        "primary_per_captured_factor_matches_three_component_formula": max_primary_captured_difference
        <= tolerance,
        "max_primary_per_influent_factor_abs_difference_vs_capture_scaled_formula": max_primary_influent_difference,
        "primary_per_influent_factor_equals_capture_scaled_formula": max_primary_influent_difference
        <= tolerance,
        "max_total_burden_identity_error": max_total_identity_error,
        "max_total_burden_identity_relative_error": max_total_identity_relative_error,
        "quarterly_total_equals_component_sum": max_total_identity_relative_error
        <= tolerance,
        "all_checks_passed": all(
            (
                (not require_legacy_regression) or max_reconstruction_difference <= tolerance,
                max_primary_captured_difference <= tolerance,
                max_primary_influent_difference <= tolerance,
                max_total_identity_relative_error <= tolerance,
            )
        ),
    }

    return _json_safe(
        {
            "schema_version": "1.1",
            "model_name": "IER treatment burden with capacity-utilization sensitivity",
            "period": f"{START_YEAR} Q1-{END_YEAR} Q4",
            "functional_unit": "1 kg PFAS entering wastewater management",
            "equation": (
                "BF_influent = m_IER*S + (m_IER + m_PFAS,captured)*D_excluding_lime + m_PFAS,captured*L; "
                "quarterly burden = wastewater PFAS input * BF_influent"
            ),
            "interpretation": (
                "The primary loading is read from IER_model_input.xlsx (0.1 g/g EPA-informed "
                "PFOS high-loading proxy, adopted for the mixture). Lower f_C increases resin "
                "production and non-lime mixture incineration. Lime scales only with captured "
                "PFAS. Legacy regression is evaluated separately at 1.05 g/g."
            ),
            "scope": (
                "Central IER production/separation and spent-mixture incineration LCI only. "
                "Legacy uncertainty/error columns, bypass toxicity, incineration residual "
                "toxicity, and other waste-media treatments are outside this stage."
            ),
            "inputs": {
                "attribution_output": attribution_source,
                "ier_parameter_workbook": ier_model_results.get("input_path", str(Path(ier_input_path).resolve())),
                "legacy_central_factor_workbook": str(legacy_path),
                "lci_2025_workbook": str(source_2025_path),
                "lci_2030_scenario_workbooks": [str(path) for path in source_2030_paths],
            },
            "reference_batch": {
                "ier_kg": REFERENCE_IER_KG,
                "captured_pfas_kg": REFERENCE_CAPTURED_PFAS_KG,
                "spent_mix_kg": REFERENCE_SPENT_MIX_KG,
                "note": (
                    "Legacy separation is reported for 1 kg IER and legacy destruction for "
                    "the associated 2.05 kg spent mixture."
                ),
            },
            "mass_model_assumptions": ier_model_results.get("assumptions", {}),
            "primary_scenario_id": primary_id,
            "impact_categories": impacts,
            "mass_scenarios": mass_scenarios,
            "lci_anchors": [
                {
                    "impact_index": impact["impact_index"],
                    "impact_code": impact["impact_code"],
                    "unit": impact["unit"],
                    "separation_2025_per_kg_ier": float(separation_2025[index]),
                    "destruction_2025_per_reference_2p05kg_batch": float(
                        destruction_2025_batch[index]
                    ),
                    "destruction_excluding_lime_2025_per_kg_spent_mix": float(destruction_no_lime_2025[index]),
                    "lime_2025_per_kg_captured_pfas": float(lime_2025[index]),
                    "separation_2030_average_per_kg_ier": float(
                        separation_2030[index]
                    ),
                    "destruction_2030_average_per_reference_2p05kg_batch": float(
                        destruction_2030_batch[index]
                    ),
                    "destruction_excluding_lime_2030_average_per_kg_spent_mix": float(destruction_no_lime_2030[index]),
                    "lime_2030_average_per_kg_captured_pfas": float(lime_2030[index]),
                }
                for index, impact in enumerate(impacts)
            ],
            "quarterly_factors": quarterly_factors,
            "quarterly_burden": quarterly_burden,
            "annual_burden": annual_burden,
            "cumulative_burden_2025_2030": cumulative_burden,
            "quality_checks": quality_checks,
        }
    )


def save_json(
    payload: dict[str, Any], output_path: str | Path = DEFAULT_OUTPUT
) -> Path:
    """Write the IER burden results and return the resolved output path."""

    resolved_output = Path(output_path).expanduser().resolve()
    resolved_output.parent.mkdir(parents=True, exist_ok=True)
    resolved_output.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return resolved_output


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build mass-consistent IER treatment burdens for 2025-2030."
    )
    parser.add_argument("--attribution", type=Path, default=DEFAULT_ATTRIBUTION)
    parser.add_argument("--ier-input", type=Path, default=DEFAULT_IER_INPUT)
    parser.add_argument("--legacy-factor", type=Path, default=DEFAULT_LEGACY_FACTOR)
    parser.add_argument("--lci-2025", type=Path, default=DEFAULT_LCI_2025)
    parser.add_argument(
        "--lci-2030",
        type=Path,
        nargs=4,
        default=list(DEFAULT_LCI_2030),
        metavar=("NDI", "PK500", "PK1150", "SSP5_BASE"),
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    ier_model_results = run_ier_model(args.ier_input)
    payload = calculate_ier_burden(
        attribution_results=args.attribution,
        ier_model_results=ier_model_results,
        ier_input_path=args.ier_input,
        legacy_factor_path=args.legacy_factor,
        lci_2025_path=args.lci_2025,
        lci_2030_paths=args.lci_2030,
    )
    output = save_json(payload, args.output)
    cc_primary = next(
        row
        for row in payload["cumulative_burden_2025_2030"]
        if row["scenario_id"] == payload["primary_scenario_id"]
        and row["impact_code"] == "CC"
    )
    print(
        json.dumps(
            {
                "output": str(output),
                "primary_scenario_id": payload["primary_scenario_id"],
                "primary_cumulative_CC": cc_primary["total_ier_treatment_burden"],
                "primary_cumulative_CC_unit": cc_primary["unit"],
                "all_checks_passed": payload["quality_checks"]["all_checks_passed"],
            },
            indent=2,
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()

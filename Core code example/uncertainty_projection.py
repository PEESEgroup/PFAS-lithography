"""Normal path-wise parameter biases with retained compute and vintage retirements."""
from pathlib import Path

import numpy as np
import xlrd
from openpyxl import load_workbook

DEFAULT_UNCERTAINTY_INPUT = Path(__file__).resolve().parents[1] / 'Data input' / 'Uncertainty_model_input.xlsx'
SCENARIO_COLUMNS = {'aggressive': (13, 14), 'moderate': (15, 16), 'conservative': (17, 18)}


def read_uncertainty_settings(path=DEFAULT_UNCERTAINTY_INPUT):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        settings = {r[0]: r[1] for r in workbook['Settings'].iter_rows(min_row=2, values_only=True) if r[0]}
    finally:
        workbook.close()
    required = {'runs', 'seed', 'ai_cv_initial', 'ai_cv_final', 'cv_flat_through',
                'cv_reaches_final', 'daily_token_cv', 'manufacturing_cv', 'distribution',
                'ai_temporal_sampling', 'daily_token_temporal_sampling',
                'manufacturing_temporal_sampling', 'cross_parameter_sampling'}
    if required - settings.keys():
        raise ValueError(f'Missing uncertainty settings: {required - settings.keys()}')
    for key, expected in {'distribution': 'normal', 'ai_temporal_sampling': 'shared_path',
                          'daily_token_temporal_sampling': 'shared_path',
                          'manufacturing_temporal_sampling': 'shared_path',
                          'cross_parameter_sampling': 'independent'}.items():
        if settings[key] != expected:
            raise ValueError(f'{key} must be {expected}; other sampling models are not implemented')
    for key in ('runs', 'seed'):
        if int(settings[key]) != settings[key]:
            raise ValueError(f'{key} must be an integer')
        settings[key] = int(settings[key])
    if settings['runs'] < 2 or settings['seed'] < 0:
        raise ValueError('Need at least two runs and a nonnegative seed')
    for key in ('ai_cv_initial', 'ai_cv_final', 'daily_token_cv', 'manufacturing_cv'):
        settings[key] = float(settings[key])
        if not np.isfinite(settings[key]) or settings[key] < 0:
            raise ValueError(f'{key} must be finite and nonnegative')
    if float(settings['cv_reaches_final']) <= float(settings['cv_flat_through']):
        raise ValueError('Final CV quarter must follow the last flat-CV quarter')
    return settings


def quarterly_cv(periods, settings):
    start, end = float(settings['cv_flat_through']), float(settings['cv_reaches_final'])
    fraction = np.clip((np.asarray(periods) - start) / (end - start), 0, 1)
    return settings['ai_cv_initial'] + fraction * (settings['ai_cv_final'] - settings['ai_cv_initial'])


def normal_samples(rng, mean, cv, size, name, *, shared_path=False):
    """Untruncated normal; report invalid physical inputs instead of silently clipping/redrawing."""
    if shared_path:
        # Each parameter has its own z per run; z is reused across all quarters.
        # Changing CV changes the deviation's magnitude, not its sampled sign.
        z = rng.standard_normal((size[0], 1))
        values = np.broadcast_to(mean, size) * (1 + np.asarray(cv) * z)
    else:
        values = rng.normal(mean, np.asarray(mean) * cv, size=size)
    if not np.all(np.isfinite(values)) or np.any(values <= 0):
        raise ValueError(f'Nonpositive/nonfinite normal draw in {name}; review distribution/bounds before rerunning')
    return values


def purchase_path(demand, performance, retirement):
    """Arrays [run, quarter]; retirement[a-1] is the original-cohort age-a outflow fraction."""
    demand, performance = np.asarray(demand), np.asarray(performance)
    if demand.shape != performance.shape or np.any(demand < 0) or np.any(performance <= 0):
        raise ValueError('Demand/performance must align, with nonnegative demand and positive performance')
    retirement = np.asarray(retirement, dtype=float)
    if np.any(retirement < 0) or not np.isclose(retirement.sum(), 1.0, atol=1e-10):
        raise ValueError('Retirement fractions must be nonnegative and sum to one')
    additions = np.zeros_like(demand, dtype=float)
    retired = np.zeros_like(additions)
    capacity = np.zeros_like(additions)
    stock = np.zeros_like(additions)
    for t in range(demand.shape[1]):
        retired_capacity = np.zeros(demand.shape[0])
        for age in range(1, min(t, len(retirement)) + 1):
            cohort_retired = additions[:, t-age] * retirement[age-1]
            retired[:, t] += cohort_retired
            retired_capacity += cohort_retired * performance[:, t-age]
        available = (capacity[:, t-1] if t else 0.0) - retired_capacity
        available_stock = (stock[:, t-1] if t else 0.0) - retired[:, t]
        if np.any(available < -1e-6) or np.any(available_stock < -1e-6):
            raise AssertionError('Retirements exceed installed inventory')
        available, available_stock = np.maximum(available, 0.0), np.maximum(available_stock, 0.0)
        additions[:, t] = np.ceil(np.maximum(demand[:, t] - available, 0.0) / performance[:, t])
        capacity[:, t] = available + additions[:, t] * performance[:, t]
        stock[:, t] = available_stock + additions[:, t]
    if np.any(capacity + 1e-6 < demand):
        raise AssertionError('Purchased compute does not meet demand')
    return {'new_servers': additions, 'retired_servers': retired,
            'installed_compute': capacity, 'idle_compute': np.maximum(capacity-demand, 0.0),
            'installed_servers': stock}


def run_server_projection(input_path='Data input/legacy_data.xls', *, uncertainty_input=DEFAULT_UNCERTAINTY_INPUT,
                          settings=None, scenario='aggressive', seed=None, deterministic=False,
                          inference_efficiency_multiplier=1.0):
    # Training and inference use the same demand equations as the study.
    from demand_equations import calc_Rt, calc_Ri
    settings = dict(settings if settings is not None else read_uncertainty_settings(uncertainty_input))
    if not np.isfinite(inference_efficiency_multiplier) or inference_efficiency_multiplier <= 0:
        raise ValueError('Inference efficiency multiplier must be finite and positive')
    if scenario not in SCENARIO_COLUMNS:
        raise ValueError(f'Unknown deployment scenario: {scenario}')
    excel = xlrd.open_workbook(str(input_path))
    config, data = excel.sheet_by_name('input Configuration'), excel.sheet_by_name('Computation Data.')
    if int(config.cell_value(4, 1)) != 1:
        raise ValueError('Uncertainty workflow supports the active lifetime-based purchase strategy only')
    def col(index):
        values = np.asarray(data.col_values(index, 1), dtype=float)
        if not np.all(np.isfinite(values)):
            raise ValueError(f'Nonfinite legacy input column {index}')
        return values
    periods = col(11)
    if not np.allclose(np.diff(periods), .25):
        raise ValueError('Expected consecutive legacy quarters')
    n = 2 if deterministic else settings['runs']
    effective_seed = settings['seed'] if seed is None else seed
    rng = np.random.default_rng(np.random.SeedSequence([effective_seed, 0]))
    cv = np.zeros(len(periods)) if deterministic else quarterly_cv(periods, settings)
    inputs = {}
    for name, column in [('nt_para', 0), ('nt_traindata', 1), ('pt_theo', 2), ('pi_theo', 3)]:
        inputs[name] = normal_samples(rng, col(column), cv, (n, len(periods)), name, shared_path=True)
    inputs['daily_tokens'] = normal_samples(rng, float(config.cell_value(5, 1)),
                                            0 if deterministic else settings['daily_token_cv'],
                                            (n, len(periods)), 'daily_tokens', shared_path=True)
    users_col, models_col = SCENARIO_COLUMNS[scenario]
    inputs['ni_user'], inputs['n_model'] = col(users_col), col(models_col)
    gpu_per_server = float(config.cell_value(3, 1))
    train = calc_Rt(inputs['nt_para'], inputs['nt_traindata'], inputs['n_model'], col(7), float(config.cell_value(2, 1)))
    # The exact same nt_para sample is reused in training and inference.
    infer = calc_Ri(inputs['nt_para'], inputs['ni_user'], inputs['daily_tokens'], col(7))
    pt, pi = inputs['pt_theo'] * col(5) * gpu_per_server, inputs['pi_theo'] * col(6) * gpu_per_server
    # Apply to every cohort over the complete initialization/reporting path.
    pi = pi * inference_efficiency_multiplier
    retirement = np.asarray(excel.sheet_by_name('Lifespan normal distribution').col_values(0, 1), dtype=float)
    padding = int(config.cell_value(1, 1)) - len(retirement)//2
    if padding < 0:
        raise ValueError('Legacy lifespan shift would truncate the retirement distribution')
    retirement = np.pad(retirement, (padding, padding))
    train_state, infer_state = purchase_path(train, pt, retirement), purchase_path(infer, pi, retirement)
    additions = train_state['new_servers'] + infer_state['new_servers']
    outflow = train_state['retired_servers'] + infer_state['retired_servers']
    requirement = np.pad(train + infer, ((0, 0), (1, 0)))
    manufacturing = normal_samples(np.random.default_rng(np.random.SeedSequence([effective_seed, 1])),
                                   1.0, 0 if deterministic else settings['manufacturing_cv'],
                                   (n,), 'manufacturing_multiplier')
    result = {'scenario': scenario, 'periods': periods, 'settings': settings, 'seed': effective_seed,
              'inference_efficiency_multiplier': float(inference_efficiency_multiplier),
              'manufacturing_multipliers': manufacturing, 'deterministic': deterministic,
              'input_samples': inputs, 'samples': additions, 'training': train_state, 'inference': infer_state,
              'gpu_per_server': gpu_per_server, 'ai_cv_by_quarter': cv,
              'training_demand': train, 'inference_demand': infer,
              'quality_checks': {'nonnegative_purchases': bool(np.all(additions >= 0)),
                                 'finite_samples': bool(np.all(np.isfinite(additions))),
                                 'retirement_share_sum': float(retirement.sum())}}
    for values, mean_key, sd_key, cumulative_key in [
            (additions, 'mean_list', 'std_list', 'std_cum_list'),
            (outflow, 'mean_out', 'std_out_list', 'std_cum_out_list'),
            (requirement, 'mean_R', 'std_R', None)]:
        result[mean_key] = values.mean(axis=0).tolist()
        result[sd_key] = values.std(axis=0, ddof=1).tolist()
        if cumulative_key:
            result[cumulative_key] = values.cumsum(axis=1).std(axis=0, ddof=1).tolist()
    return result

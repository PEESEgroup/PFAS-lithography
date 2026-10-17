"""Calculate gas, solvent, and solid-waste PFAS incineration burdens.

The active manufacturing model passes kilograms of PFAS assigned to each waste
medium.  This module preserves the legacy treatment factors and their original
per-kilogram-PFAS screening interpretation while replacing the legacy tonnes to
kilograms interface.  Ordinary hazardous-waste incineration is used for gas and
solid PFAS; the separate high-energy solvent inventory is used for solvent PFAS.

The process burden is scaled by the full PFAS mass sent to incineration.  Parent-
PFAS DRE is intentionally absent here because the toxicity module already applies
the DRE once when calculating residual parent PFAS released to air.
"""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import openpyxl
import xlrd

from attribution.genai_pfas_attribution import MATERIAL_NAMES
from ier.ier_burden import IMPACT_CODES


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_ATTRIBUTION = (
    PROJECT_ROOT / "Excel results" / "attribution" / "GenAI_PFAS_attribution_results.json"
)
DEFAULT_LEGACY_INPUT = PROJECT_ROOT / "Data input" / "legacy_data.xls"
DEFAULT_LEGACY_FACTORS = (
    PROJECT_ROOT / "Data input" / "treatment" / "treatment_destruction_summary_new.xls"
)
DEFAULT_LCI_2025 = (
    PROJECT_ROOT
    / "Data input"
    / "treatment"
    / "treatment_Database-IAM_scenarios_Basic_2025.xlsx"
)
DEFAULT_LCI_2030 = (
    PROJECT_ROOT / "Data input" / "treatment" / "treatment_PFAS_destruction_NDI_2030.xlsx",
    PROJECT_ROOT
    / "Data input"
    / "treatment"
    / "treatment_PFAS_destruction_PkDuge500_2030.xlsx",
    PROJECT_ROOT
    / "Data input"
    / "treatment"
    / "treatment_PFAS_destruction_PkDuge1150_2030.xlsx",
    PROJECT_ROOT
    / "Data input"
    / "treatment"
    / "treatment_PFAS_destruction_SSP5-BASE_2030.xlsx",
)
DEFAULT_OUTPUT = (
    PROJECT_ROOT / "Excel results" / "treatment" / "Incineration_burden_results.json"
)

INCINERATED_MEDIA = ("gas", "solvent", "solid")
ORDINARY_MEDIA = {"gas", "solid"}
IMPACT_COUNT = len(IMPACT_CODES)
EXPECTED_PERIODS = tuple(
    f"{year} Q{quarter}"
    for year in range(2025, 2031)
    for quarter in range(1, 5)
)


class IncinerationInputError(ValueError):
    """Raised when an incineration source violates the expected contract."""


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return None if not math.isfinite(float(value)) else float(value)
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    return value


def _read_json(source: str | Path | dict[str, Any]) -> tuple[dict[str, Any], str]:
    if isinstance(source, dict):
        return source, "in-memory active attribution output"
    path = Path(source).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"Attribution output not found: {path}")
    return json.loads(path.read_text(encoding="utf-8")), str(path)


def _numeric_vector(values: Iterable[Any], context: str) -> np.ndarray:
    try:
        result = np.asarray(list(values), dtype=float)
    except (TypeError, ValueError) as exc:
        raise IncinerationInputError(f"{context} contains non-numeric values") from exc
    if result.shape != (IMPACT_COUNT,) or not np.isfinite(result).all():
        raise IncinerationInputError(
            f"{context} must contain {IMPACT_COUNT} finite impact values"
        )
    return result


def _find_sheet(workbook: openpyxl.Workbook, names: Iterable[str], path: Path):
    for name in names:
        if name in workbook.sheetnames:
            return workbook[name]
    raise IncinerationInputError(
        f"{path.name} is missing all accepted sheets: {list(names)}"
    )


def _read_lci_factor(
    path: Path, sheet_names: Iterable[str]
) -> tuple[np.ndarray, list[str]]:
    if not path.exists():
        raise FileNotFoundError(f"Incineration LCI workbook not found: {path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = _find_sheet(workbook, sheet_names, path)
        label = str(sheet.cell(31, 2).value or "").strip().lower()
        unit = str(sheet.cell(31, 3).value or "").strip().lower()
        amount = float(sheet.cell(31, 4).value)
        if "hazardous" not in label or unit != "kg" or not math.isclose(amount, 1.0):
            raise IncinerationInputError(
                f"{path.name}/{sheet.title} row 31 is not the expected 1 kg result"
            )
        factors = _numeric_vector(
            (sheet.cell(31, column).value for column in range(5, 24)),
            f"{path.name}/{sheet.title} row 31",
        )
        methods = [
            str(sheet.cell(2, column).value or "").strip()
            for column in range(5, 24)
        ]
    finally:
        workbook.close()
    return factors, methods


def _read_legacy_factor_block(
    path: Path, sheet_name: str
) -> tuple[list[str], list[str], np.ndarray]:
    if not path.exists():
        raise FileNotFoundError(f"Legacy treatment workbook not found: {path}")
    workbook = xlrd.open_workbook(str(path))
    if sheet_name not in workbook.sheet_names() or "2030" not in workbook.sheet_names():
        raise IncinerationInputError(
            f"{path.name} is missing {sheet_name!r} or the unit sheet '2030'"
        )
    sheet = workbook.sheet_by_name(sheet_name)
    unit_sheet = workbook.sheet_by_name("2030")
    headers = [str(value).strip() for value in sheet.row_values(0, 1, 20)]
    units = [str(value).strip() for value in unit_sheet.row_values(1, 1, 20)]
    rows: list[list[float]] = []
    periods: list[str] = []
    for row_index in range(1, sheet.nrows):
        decimal_year = sheet.cell_value(row_index, 0)
        if not isinstance(decimal_year, (int, float)):
            continue
        if 2025 - 1e-9 <= decimal_year <= 2030.75 + 1e-9:
            year = int(math.floor(decimal_year + 1e-9))
            quarter = int(round((decimal_year - year) * 4)) + 1
            periods.append(f"{year} Q{quarter}")
            rows.append(sheet.row_values(row_index, 1, 20))
    values = np.asarray(rows, dtype=float)
    if tuple(periods) != EXPECTED_PERIODS or values.shape != (
        len(EXPECTED_PERIODS),
        IMPACT_COUNT,
    ):
        raise IncinerationInputError(
            f"{sheet_name} must contain 24 quarters by 19 impacts"
        )
    return headers, units, values


def _read_emission_factors(path: Path) -> np.ndarray:
    if not path.exists():
        raise FileNotFoundError(f"Legacy emission-factor workbook not found: {path}")
    workbook = xlrd.open_workbook(str(path))
    if "Emission factor" not in workbook.sheet_names():
        raise IncinerationInputError(f"{path.name} is missing 'Emission factor'")
    sheet = workbook.sheet_by_name("Emission factor")
    values = np.asarray(
        [sheet.row_values(row, 1, 7) for row in range(1, 5)], dtype=float
    )
    if values.shape != (4, len(MATERIAL_NAMES)) or not np.isfinite(values).all():
        raise IncinerationInputError("Emission-factor matrix must be four by six")
    return values


def _aggregate(
    rows: list[dict[str, Any]], group_keys: tuple[str, ...]
) -> list[dict[str, Any]]:
    grouped: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        key = tuple(row[name] for name in group_keys)
        if key not in grouped:
            grouped[key] = {name: row[name] for name in group_keys}
            grouped[key].update(
                {
                    "pfas_sent_to_incineration_kg": 0.0,
                    "process_burden": 0.0,
                }
            )
            if "impact_code" in row:
                grouped[key]["unit"] = row["unit"]
        grouped[key]["pfas_sent_to_incineration_kg"] += float(
            row["pfas_sent_to_incineration_kg"]
        )
        grouped[key]["process_burden"] += float(row["process_burden"])
    return list(grouped.values())


def calculate_incineration_burden(
    attribution_results: str | Path | dict[str, Any] = DEFAULT_ATTRIBUTION,
    legacy_input_path: str | Path = DEFAULT_LEGACY_INPUT,
    legacy_factor_path: str | Path = DEFAULT_LEGACY_FACTORS,
    lci_2025_path: str | Path = DEFAULT_LCI_2025,
    lci_2030_paths: Iterable[str | Path] = DEFAULT_LCI_2030,
    *, require_legacy_regression: bool = True,
) -> dict[str, Any]:
    """Return central process burdens for gas, solvent, and solid PFAS."""

    attribution, attribution_source = _read_json(attribution_results)
    quarterly_attribution = attribution.get("quarterly", [])
    if len(quarterly_attribution) != len(EXPECTED_PERIODS):
        raise IncinerationInputError("Attribution must contain 24 quarterly rows")
    if tuple(str(row.get("Period")) for row in quarterly_attribution) != EXPECTED_PERIODS:
        raise IncinerationInputError("Attribution periods are not 2025 Q1-2030 Q4")

    source_2025 = Path(lci_2025_path).expanduser().resolve()
    source_2030 = [Path(path).expanduser().resolve() for path in lci_2030_paths]
    if len(source_2030) != 4:
        raise IncinerationInputError("Exactly four 2030 IAM workbooks are required")

    ordinary_2025, ordinary_methods_2025 = _read_lci_factor(
        source_2025, ("Incineration",)
    )
    solvent_2025, solvent_methods_2025 = _read_lci_factor(
        source_2025, ("Incineration-solvent", "Incineration solvent")
    )
    ordinary_2030_rows = [
        _read_lci_factor(path, ("Incineration",)) for path in source_2030
    ]
    solvent_2030_rows = [
        _read_lci_factor(path, ("Incineration solvent", "Incineration-solvent"))
        for path in source_2030
    ]
    ordinary_2030 = np.mean(
        np.stack([row[0] for row in ordinary_2030_rows]), axis=0
    )
    solvent_2030 = np.mean(
        np.stack([row[0] for row in solvent_2030_rows]), axis=0
    )
    ordinary_methods = ordinary_2030_rows[0][1]
    solvent_methods = solvent_2030_rows[0][1]
    if any(row[1] != ordinary_methods for row in ordinary_2030_rows[1:]):
        raise IncinerationInputError("Ordinary 2030 workbooks use different methods")
    if any(row[1] != solvent_methods for row in solvent_2030_rows[1:]):
        raise IncinerationInputError("Solvent 2030 workbooks use different methods")
    if ordinary_methods_2025 != solvent_methods_2025:
        raise IncinerationInputError("2025 ordinary and solvent methods differ")

    legacy_path = Path(legacy_factor_path).expanduser().resolve()
    legacy_headers, units, legacy_ordinary = _read_legacy_factor_block(
        legacy_path, "data for python-incineration"
    )
    _, _, legacy_solvent = _read_legacy_factor_block(
        legacy_path, "data incin solvent"
    )
    factors: dict[str, np.ndarray] = {}
    for medium in INCINERATED_MEDIA:
        start = ordinary_2025 if medium in ORDINARY_MEDIA else solvent_2025
        end = ordinary_2030 if medium in ORDINARY_MEDIA else solvent_2030
        factors[medium] = np.stack(
            [
                start + (end - start) * (quarter_index / 23)
                for quarter_index in range(24)
            ]
        )

    max_ordinary_regression_error = float(
        np.max(np.abs(factors["gas"] - legacy_ordinary))
    )
    max_solvent_regression_error = float(
        np.max(np.abs(factors["solvent"] - legacy_solvent))
    )
    emission_factors = _read_emission_factors(
        Path(legacy_input_path).expanduser().resolve()
    )
    emission_row_index = {"gas": 0, "solvent": 1, "wastewater": 2, "solid": 3}

    impacts: list[dict[str, Any]] = []
    for index, code in enumerate(IMPACT_CODES):
        unit = units[index]
        if not unit and code in {"ECF_inorganics", "ECF_organics"}:
            unit = "CTUe"
        elif not unit and code == "ER_nonrenewable":
            unit = "source unit not retained in legacy summary"
        impacts.append(
            {
                "impact_index": index + 1,
                "impact_code": code,
                "legacy_header": legacy_headers[index],
                "method_name": ordinary_methods[index],
                "unit": unit,
            }
        )

    factor_rows: list[dict[str, Any]] = []
    burden_rows: list[dict[str, Any]] = []
    material_quarter_rows: list[dict[str, Any]] = []
    max_stream_mass_error = 0.0
    max_material_media_mass_error = 0.0
    for quarter_index, attribution_row in enumerate(quarterly_attribution):
        period = EXPECTED_PERIODS[quarter_index]
        year = int(attribution_row["Year"])
        quarter = int(attribution_row["Quarter"])
        material_masses = np.asarray(
            [
                float(attribution_row[f"Material | {material} (kg)"])
                for material in MATERIAL_NAMES
            ],
            dtype=float,
        )
        for medium in INCINERATED_MEDIA:
            stream_mass = float(
                attribution_row[f"Manufacturing flow | {medium} (kg)"]
            )
            medium_material_mass = (
                material_masses * emission_factors[emission_row_index[medium], :]
            )
            max_material_media_mass_error = max(
                max_material_media_mass_error,
                abs(float(medium_material_mass.sum()) - stream_mass),
            )
            for impact_index, impact in enumerate(impacts):
                factor = float(factors[medium][quarter_index, impact_index])
                burden = stream_mass * factor
                factor_rows.append(
                    {
                        "period": period,
                        "year": year,
                        "quarter": quarter,
                        "medium": medium,
                        "inventory_proxy": (
                            "ordinary hazardous-waste incineration"
                            if medium in ORDINARY_MEDIA
                            else "high-energy PFAS solvent incineration"
                        ),
                        "impact_code": impact["impact_code"],
                        "unit": impact["unit"],
                        "factor_per_kg_pfas_sent": factor,
                    }
                )
                burden_rows.append(
                    {
                        "period": period,
                        "year": year,
                        "quarter": quarter,
                        "medium": medium,
                        "impact_code": impact["impact_code"],
                        "unit": impact["unit"],
                        "pfas_sent_to_incineration_kg": stream_mass,
                        "factor_per_kg_pfas_sent": factor,
                        "process_burden": burden,
                    }
                )
                material_sum = 0.0
                for material_index, material in enumerate(MATERIAL_NAMES):
                    material_mass = float(medium_material_mass[material_index])
                    material_burden = material_mass * factor
                    material_sum += material_burden
                    material_quarter_rows.append(
                        {
                            "period": period,
                            "year": year,
                            "quarter": quarter,
                            "medium": medium,
                            "material": material,
                            "impact_code": impact["impact_code"],
                            "unit": impact["unit"],
                            "pfas_sent_to_incineration_kg": material_mass,
                            "process_burden": material_burden,
                        }
                    )
                max_stream_mass_error = max(
                    max_stream_mass_error, abs(material_sum - burden)
                )

    annual_by_medium = _aggregate(
        burden_rows, ("year", "medium", "impact_code")
    )
    cumulative_by_medium = _aggregate(
        burden_rows, ("medium", "impact_code")
    )
    annual_total = _aggregate(burden_rows, ("year", "impact_code"))
    cumulative_total = _aggregate(burden_rows, ("impact_code",))
    annual_by_material = _aggregate(
        material_quarter_rows, ("year", "medium", "material", "impact_code")
    )
    cumulative_by_material = _aggregate(
        material_quarter_rows, ("medium", "material", "impact_code")
    )

    max_annual_cumulative_error = 0.0
    max_medium_total_error = 0.0
    max_material_medium_error = 0.0
    for medium_row in cumulative_by_medium:
        annual_rows = [
            row
            for row in annual_by_medium
            if row["medium"] == medium_row["medium"]
            and row["impact_code"] == medium_row["impact_code"]
        ]
        material_rows = [
            row
            for row in cumulative_by_material
            if row["medium"] == medium_row["medium"]
            and row["impact_code"] == medium_row["impact_code"]
        ]
        for field in ("pfas_sent_to_incineration_kg", "process_burden"):
            max_annual_cumulative_error = max(
                max_annual_cumulative_error,
                abs(
                    float(medium_row[field])
                    - sum(float(row[field]) for row in annual_rows)
                ),
            )
            max_material_medium_error = max(
                max_material_medium_error,
                abs(
                    float(medium_row[field])
                    - sum(float(row[field]) for row in material_rows)
                ),
            )
    for total_row in cumulative_total:
        medium_rows = [
            row
            for row in cumulative_by_medium
            if row["impact_code"] == total_row["impact_code"]
        ]
        for field in ("pfas_sent_to_incineration_kg", "process_burden"):
            max_medium_total_error = max(
                max_medium_total_error,
                abs(
                    float(total_row[field])
                    - sum(float(row[field]) for row in medium_rows)
                ),
            )

    tolerance = 1e-9
    quality_checks = {
        "quarter_count": len(quarterly_attribution),
        "medium_count": len(INCINERATED_MEDIA),
        "impact_count": IMPACT_COUNT,
        "max_legacy_ordinary_factor_difference": max_ordinary_regression_error,
        "max_legacy_solvent_factor_difference": max_solvent_regression_error,
        "legacy_regression_required": require_legacy_regression,
        "max_material_to_stream_mass_error_kg": max_material_media_mass_error,
        "max_material_to_stream_burden_error": max_stream_mass_error,
        "max_annual_to_cumulative_reconciliation_error": max_annual_cumulative_error,
        "max_medium_to_total_reconciliation_error": max_medium_total_error,
        "max_material_to_medium_reconciliation_error": max_material_medium_error,
        "legacy_tonnes_to_kg_multiplier_removed": True,
        "dre_applied_in_process_burden": False,
    }
    quality_checks["all_checks_passed"] = all(
        (
            (not require_legacy_regression) or max_ordinary_regression_error <= 1e-10,
            (not require_legacy_regression) or max_solvent_regression_error <= 1e-10,
            max_material_media_mass_error <= tolerance,
            max_stream_mass_error <= tolerance,
            max_annual_cumulative_error <= tolerance,
            max_medium_total_error <= tolerance,
            max_material_medium_error <= tolerance,
        )
    )

    return _json_safe(
        {
            "schema_version": "1.0",
            "model_name": "Gas, solvent, and solid PFAS incineration process burden",
            "period": "2025 Q1-2030 Q4",
            "functional_unit": (
                "1 kg PFAS assigned to the treated gas, solvent, or solid-waste stream"
            ),
            "attribution_interface": attribution_source,
            "mass_basis": (
                "Full PFAS mass sent to incineration. Factors retain the legacy "
                "per-kg-PFAS screening basis; they are not per kg of total carrier waste."
            ),
            "dre_boundary": (
                "DRE does not scale process burden; the toxicity module applies it once "
                "to residual parent PFAS."
            ),
            "impacts": impacts,
            "factor_anchors": {
                "ordinary_2025": ordinary_2025.tolist(),
                "ordinary_2030_four_scenario_mean": ordinary_2030.tolist(),
                "solvent_2025_center": solvent_2025.tolist(),
                "solvent_2030_four_scenario_mean": solvent_2030.tolist(),
            },
            "quarterly_factors": factor_rows,
            "quarterly_burden": burden_rows,
            "annual_by_medium": annual_by_medium,
            "cumulative_by_medium": cumulative_by_medium,
            "annual_total": annual_total,
            "cumulative_total": cumulative_total,
            "annual_by_material": annual_by_material,
            "cumulative_by_material": cumulative_by_material,
            "sources": {
                "legacy_factors": str(legacy_path),
                "lci_2025": str(source_2025),
                "lci_2030": [str(path) for path in source_2030],
                "emission_factors": str(Path(legacy_input_path).expanduser().resolve()),
            },
            "limitations": [
                "Gas and solid PFAS use the same ordinary hazardous-waste incineration proxy.",
                "The solvent center is read from the explicit local solvent inventory; heat assumptions are verified by the strategy entry point.",
                "Carrier-waste mass, PIC/PID formation, and treatment adoption shares are outside this module.",
            ],
            "quality_checks": quality_checks,
        }
    )


def save_json(payload: dict[str, Any], output_path: str | Path = DEFAULT_OUTPUT) -> Path:
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attribution", type=Path, default=DEFAULT_ATTRIBUTION)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    payload = calculate_incineration_burden(args.attribution)
    if not payload["quality_checks"]["all_checks_passed"]:
        raise RuntimeError("Incineration burden checks failed")
    output = save_json(payload, args.output)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
